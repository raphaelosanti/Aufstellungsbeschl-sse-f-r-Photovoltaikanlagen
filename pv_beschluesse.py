#!/usr/bin/env python3
"""
PV-Beschlussregister (MVP, Pilot Brandenburg)
=============================================

Sammelt Beschlüsse zu Photovoltaik-Bebauungsplänen aus Ratsinformationssystemen
(OParl-API oder RIS-HTML), extrahiert Projektdaten (heuristisch oder per Gemini)
und schreibt eine Excel-Datei mit den Tabellen "Projekte" und "Beschluesse".

Aufruf:
    python pv_beschluesse.py --quellen quellen_brandenburg.csv --seit 2023-01-01
    python pv_beschluesse.py --quellen quellen_brandenburg.csv --llm gemini   # GEMINI_API_KEY setzen

Abhängigkeiten: pip install requests beautifulsoup4 openpyxl pypdf
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import os
import re
import sys
import time
import urllib.robotparser
from dataclasses import dataclass, field, asdict
from difflib import SequenceMatcher
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

USER_AGENT = ("PV-Beschlussregister-MVP/0.1 (Marktbeobachtung; Kontakt: "
              + os.environ.get("PV_KONTAKT", "bitte eintragen") + ")")
REQUEST_PAUSE_S = 1.0          # Höflichkeitspause zwischen Requests pro Host
CACHE_DIR = ".cache_pv"
RUHEND_NACH_MONATEN = 18       # kein neuer Beschluss seit X Monaten -> "Ruhend"

# --------------------------------------------------------------------------- #
# Keywords
# --------------------------------------------------------------------------- #
PV_RE = re.compile(
    r"photovoltai|solar|\bpv\b|pv-|agri-?pv|freifl[äa]chen-?anlage|sonnenenergie|energiepark",
    re.I,
)
PLANUNG_RE = re.compile(
    r"bebauungsplan|b-plan|vb-plan|vorhabenbezogen|aufstellung|fl[äa]chennutzungsplan|"
    r"satzung|bauleitplan|geltungsbereich|sondergebiet",
    re.I,
)

# Reihenfolge = Fortschritt im Verfahren
BESCHLUSS_TYPEN = [
    ("Aufstellungsbeschluss", r"aufstellungsbeschluss|aufstellung\b.{0,40}(bebauungsplan|b-plan)|einleitung (eines|des) (aufstellungs)?verfahren"),
    ("Frühzeitige Beteiligung", r"fr[üu]hzeitige (öffentlichkeits)?beteiligung|§ ?3 abs\.? ?1|§ ?4 abs\.? ?1"),
    ("Entwurfs-/Auslegungsbeschluss", r"auslegungsbeschluss|offenlage|entwurfsbeschluss|billigung des entwurfs|öffentliche auslegung|veröffentlichung im internet|§ ?3 abs\.? ?2"),
    ("Abwägungsbeschluss", r"abwägung|abwaegung"),
    ("Satzungsbeschluss", r"satzungsbeschluss|als satzung|§ ?10 abs\.? ?1"),
    ("Inkrafttreten/Bekanntmachung", r"inkrafttreten|in kraft getreten|rechtskraft|bekanntmachung des satzungsbeschlusses"),
]
NEGATIV_RE = re.compile(r"abgelehnt|ablehnung|aufhebung des aufstellungsbeschlusses|einstellung des verfahrens|nicht zugestimmt", re.I)

KONZEPT_RE = [
    ("Agri-PV", r"agri-?(pv|photovoltaik)|agrophotovoltaik|landwirtschaftliche (doppel|mit)nutzung"),
    ("Moor-PV", r"moor-?(pv|photovoltaik)|wiedervernässung|moorboden"),
    ("Floating-PV", r"floating|schwimmend"),
    ("Dach-PV", r"dachfläche|aufdach|dach-pv"),
    ("Freifläche", r"freifl[äa]chen|ff-?pva|pv-?ffa|solarpark|sondergebiet (photovoltaik|solar)"),
]

FIRMA_RE = re.compile(
    r"([A-ZÄÖÜ0-9][\w&.,'\-]*(?:\s+[A-ZÄÖÜ0-9&][\w&.,'\-]*){0,5}\s+"
    r"(?:GmbH\s*&\s*Co\.?\s*KG|GmbH|AG|KG|UG(?:\s*\(haftungsbeschränkt\))?|SE|eG|mbH))"
)
FLAECHE_RE = re.compile(r"(\d{1,4}(?:[.,]\d{1,2})?)\s*(?:ha\b|hektar)", re.I)
LEISTUNG_RE = re.compile(r"(\d{1,4}(?:[.,]\d{1,2})?)\s*(MWp|MW\b|Megawatt)", re.I)
PLANNR_RE = re.compile(r"(?:Nr\.|Nummer)\s*([0-9]{1,4}[a-zA-Z/\-0-9]*)")
PLANNAME_RE = re.compile(r"[„\"“]([^“”\"]{4,120})[“”\"]")
GEMEINDE_RE = re.compile(r"\b(?:der |die )?(Gemeinde|Stadt|Ortsgemeinde|Gemeindevertretung|Stadtverordnetenversammlung)\s+([A-ZÄÖÜ][\wäöüß\-]+(?:\s(?:an der|am|im|bei)\s[A-ZÄÖÜ][\wäöüß\-]+|\s\([^)]+\))?)")
DATUM_RE = re.compile(r"\b(\d{1,2})\.(\d{1,2})\.(\d{4})\b")


# --------------------------------------------------------------------------- #
# Datenmodell
# --------------------------------------------------------------------------- #
@dataclass
class Quelle:
    name: str           # z. B. "Amt Barnim-Oderbruch"
    landkreis: str
    typ: str            # "oparl" | "html" | "auto"
    url: str            # OParl-System-URL, RIS-Basis-URL oder Vorlagenübersicht
    bundesland: str = "Brandenburg"


@dataclass
class Beschluss:
    quelle: str
    landkreis: str
    gemeinde: str
    datum: str                  # ISO yyyy-mm-dd oder ""
    titel: str
    beschlusstyp: str
    negativ: bool
    planname: str
    plannr: str
    konzept: str
    flaeche_ha: float | None
    leistung_mw: float | None
    entwickler: str
    link: str
    dokument_link: str = ""
    extraktion: str = "heuristik"   # heuristik | gemini
    projekt_id: str = ""
    bundesland: str = "Brandenburg"


# --------------------------------------------------------------------------- #
# HTTP mit Robots-Check, Pause und Cache
# --------------------------------------------------------------------------- #
class Http:
    def __init__(self, pause: float = REQUEST_PAUSE_S, cache_dir: str = CACHE_DIR):
        self.s = requests.Session()
        self.s.headers["User-Agent"] = USER_AGENT
        self.pause = pause
        self.last: dict[str, float] = {}
        self.robots: dict[str, urllib.robotparser.RobotFileParser | None] = {}
        self.cache_dir = cache_dir
        os.makedirs(cache_dir, exist_ok=True)
        self.diag: dict[str, int] = {}

    def zaehle(self, art: str) -> None:
        self.diag[art] = self.diag.get(art, 0) + 1

    def diag_reset(self) -> dict[str, int]:
        d, self.diag = self.diag, {}
        return d

    def _erlaubt(self, url: str) -> bool:
        p = urlparse(url)
        root = f"{p.scheme}://{p.netloc}"
        if root not in self.robots:
            rp = urllib.robotparser.RobotFileParser()
            try:
                r = self.s.get(root + "/robots.txt", timeout=15)
                if r.status_code == 200:
                    rp.parse(r.text.splitlines())
                else:
                    rp = None  # keine robots.txt -> erlaubt
            except requests.RequestException:
                rp = None
            self.robots[root] = rp
        rp = self.robots[root]
        return True if rp is None else rp.can_fetch(USER_AGENT, url)

    def get(self, url: str, binary: bool = False, cache: bool = True, **kw):
        if not self._erlaubt(url):
            self.zaehle("robots.txt verbietet")
            raise PermissionError(f"robots.txt verbietet: {url}")
        key = hashlib.sha1((url + json.dumps(kw.get("params", {}), sort_keys=True)).encode()).hexdigest()
        path = os.path.join(self.cache_dir, key)
        if cache and os.path.exists(path):
            with open(path, "rb") as f:
                data = f.read()
            return data if binary else data.decode("utf-8", errors="replace")
        host = urlparse(url).netloc
        wait = self.pause - (time.time() - self.last.get(host, 0))
        if wait > 0:
            time.sleep(wait)
        try:
            r = self.s.get(url, timeout=30, **kw)
        except requests.RequestException:
            self.zaehle("Verbindungsfehler")
            raise
        self.last[host] = time.time()
        if r.status_code >= 400:
            self.zaehle(f"HTTP {r.status_code}")
        r.raise_for_status()
        self.zaehle("Seiten geladen")
        if cache:
            with open(path, "wb") as f:
                f.write(r.content)
        if binary:
            return r.content
        r.encoding = r.encoding or r.apparent_encoding
        return r.text

    def json(self, url: str, **kw):
        return json.loads(self.get(url, **kw))


# --------------------------------------------------------------------------- #
# Text aus Dokumenten
# --------------------------------------------------------------------------- #
def pdf_text(data: bytes, max_seiten: int = 15) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:max_seiten])
    except Exception as e:  # noqa: BLE001
        return f"[PDF nicht lesbar: {e}]"


def html_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "nav", "header", "footer"]):
        t.decompose()
    return re.sub(r"\n{2,}", "\n", soup.get_text("\n", strip=True))


def dokument_text(http: Http, url: str) -> str:
    try:
        data = http.get(url, binary=True)
    except Exception as e:  # noqa: BLE001
        return f"[Download fehlgeschlagen: {e}]"
    if data[:4] == b"%PDF":
        return pdf_text(data)
    return html_text(data.decode("utf-8", errors="replace"))


# --------------------------------------------------------------------------- #
# Adapter 1: OParl
# --------------------------------------------------------------------------- #
OPARL_PFADE = ["oparl/1.0/system.asp", "oparl/1.1/system.asp", "oparl/system",
               "oparl/v1/system", "oparl/v1.1/system", "webservice/oparl/v1.1/system"]


def oparl_entdecken(http: Http, basis: str) -> str | None:
    """Probiert gängige OParl-Pfade relativ zur RIS-Basis-URL."""
    basis = basis if basis.endswith("/") else basis + "/"
    for pfad in OPARL_PFADE:
        url = urljoin(basis, pfad)
        try:
            d = http.json(url, cache=False)
            if isinstance(d, dict) and "body" in d and "oparl" in str(d.get("type", "")).lower():
                return url
        except Exception:  # noqa: BLE001
            continue
    return None


def oparl_liste(http: Http, url: str, params: dict | None = None, max_seiten: int = 200):
    seite = 0
    while url and seite < max_seiten:
        d = http.json(url, params=params or {}, cache=False)
        params = None  # next-Links enthalten die Parameter bereits
        for obj in d.get("data", []):
            yield obj
        url = (d.get("links") or {}).get("next")
        seite += 1


def oparl_vorlagen(http: Http, q: Quelle, seit: str | None):
    system = http.json(q.url, cache=False)
    bodies = system["body"]
    bodies = list(oparl_liste(http, bodies)) if isinstance(bodies, str) else bodies
    for body in bodies:
        body = http.json(body) if isinstance(body, str) else body
        paper_url = body.get("paper")
        if not paper_url:
            continue
        params = {"modified_since": seit + "T00:00:00+00:00"} if seit else None
        try:
            papers = list(oparl_liste(http, paper_url, params))
        except requests.HTTPError:
            papers = list(oparl_liste(http, paper_url))  # Server kennt Filter nicht
        for p in papers:
            titel = p.get("name") or ""
            if not PV_RE.search(titel):
                continue
            datum = (p.get("date") or "")[:10]
            if seit and datum and datum < seit:
                continue
            main = p.get("mainFile") or {}
            if isinstance(main, str):
                try:
                    main = http.json(main)
                except Exception:  # noqa: BLE001
                    main = {}
            doc_url = main.get("accessUrl") or main.get("downloadUrl") or ""
            text = main.get("text") or (dokument_text(http, doc_url) if doc_url else "")
            yield {
                "titel": titel,
                "datum": datum,
                "text": f"{titel}\n{p.get('reference', '')}\n{text}",
                "link": p.get("web") or p.get("id", ""),
                "dokument_link": doc_url,
                "body_name": body.get("name") or body.get("shortName") or q.name,
            }


# --------------------------------------------------------------------------- #
# Adapter 2: RIS-HTML (SessionNet *.php, ALLRIS net *.asp) – best effort
# --------------------------------------------------------------------------- #
DETAIL_LINK_RE = re.compile(r"(vo0050\.php\?__kvonr=\d+|vo020\.asp\?VOLFDNR=\d+|vo020\?VOLFDNR=\d+)", re.I)
UEBERSICHT_PFADE = ["vo0040.php", "vo040.asp", "bi/vo0040.php", "bi/vo040.asp"]


def html_vorlagen(http: Http, q: Quelle, seit: str | None, max_seiten: int = 30):
    start_urls = [q.url] if re.search(r"vo0?040", q.url) else \
        [urljoin(q.url if q.url.endswith("/") else q.url + "/", p) for p in UEBERSICHT_PFADE]
    gesehen_seiten, gesehen_detail = set(), set()
    queue = list(start_urls)
    while queue and len(gesehen_seiten) < max_seiten:
        url = queue.pop(0)
        if url in gesehen_seiten:
            continue
        gesehen_seiten.add(url)
        try:
            html = http.get(url, cache=False)
        except Exception:  # noqa: BLE001
            continue
        soup = BeautifulSoup(html, "html.parser")
        for a in soup.find_all("a", href=True):
            href, txt = a["href"], a.get_text(" ", strip=True)
            absu = urljoin(url, href)
            if DETAIL_LINK_RE.search(href):
                zeile = a.find_parent("tr")
                kontext = zeile.get_text(" ", strip=True) if zeile else txt
                if absu in gesehen_detail or not PV_RE.search(kontext):
                    continue
                gesehen_detail.add(absu)
                try:
                    detail = html_text(http.get(absu))
                except Exception:  # noqa: BLE001
                    continue
                datum = erstes_datum(detail)
                if seit and datum and datum < seit:
                    continue
                # PDF-Anhänge der Vorlage (Beschlussvorlage) mitnehmen
                pdfs = [urljoin(absu, x["href"]) for x in BeautifulSoup(http.get(absu), "html.parser")
                        .find_all("a", href=True) if re.search(r"\.pdf|getfile|do0?0?\d", x["href"], re.I)][:2]
                anhang = "\n".join(dokument_text(http, u)[:15000] for u in pdfs)
                yield {"titel": kontext[:300], "datum": datum, "text": detail + "\n" + anhang,
                       "link": absu, "dokument_link": pdfs[0] if pdfs else "", "body_name": q.name}
            elif re.search(r"(weiter|nächste|next|›|»)", txt, re.I) and re.search(r"vo0?040", absu):
                queue.append(absu)


# --------------------------------------------------------------------------- #
# Adapter 3: SessionNet (Somacos) – läuft über die Sitzungen
# --------------------------------------------------------------------------- #
# SessionNet 5.x hat oft keine öffentliche Vorlagenliste. Sitzungen haben aber
# fortlaufende Nummern (si0057.php?__ksinr=N) über alle Gemeinden des Amtes.
# Wir starten bei der höchsten bekannten Nummer und gehen rückwärts, bis die
# Sitzungen älter als --seit sind.
SI_ID_RE = re.compile(r"(?:__ksinr=|rssgo\.php\?si)(\d+)")
TOP_LINK_RE = re.compile(r"(to0050\.(?:php|asp)\?__ktonr=\d+|vo0050\.(?:php|asp)\?__kvonr=\d+)", re.I)


def _basis(url: str) -> str:
    p = urlparse(url)
    pfad = p.path.rsplit("/", 1)[0] + "/" if not p.path.endswith("/") else p.path
    return f"{p.scheme}://{p.netloc}{pfad}"


def ist_sessionnet(http: Http, basis: str) -> bool:
    for pfad in ("info.php", "info.asp", "gr0040.php", "gr0040.asp", ""):
        try:
            if "sessionnet" in http.get(urljoin(basis, pfad), cache=False).lower():
                return True
        except Exception:  # noqa: BLE001
            continue
    return False


def sessionnet_endung(http: Http, basis: str) -> str:
    """SessionNet gibt es als PHP- und als ASP-Variante."""
    for ext in ("php", "asp"):
        try:
            if "sessionnet" in http.get(urljoin(basis, f"info.{ext}"), cache=False).lower():
                return ext
        except Exception:  # noqa: BLE001
            continue
    return "php"


def sessionnet_hoechste_id(http: Http, basis: str, ext: str = "php") -> int:
    ids = []
    for pfad in (f"rssfeed.{ext}", f"si0040.{ext}", f"gr0040.{ext}", f"info.{ext}"):
        try:
            ids += [int(x) for x in SI_ID_RE.findall(http.get(urljoin(basis, pfad), cache=False))]
        except Exception:  # noqa: BLE001
            continue
    return max(ids) if ids else 0


def sitzungsdatum(text: str) -> str:
    m = re.search(r"(?:Datum|Sitzung am|vom)\s*:?\s*\w*,?\s*(\d{1,2}\.\d{1,2}\.\d{4})", text)
    return erstes_datum(m.group(1)) if m else erstes_datum(text)


def sessionnet_vorlagen(http: Http, q: Quelle, seit: str | None,
                        max_sitzungen: int = 2500, stop_nach_alt: int = 60):
    basis = _basis(q.url)
    ext = sessionnet_endung(http, basis)
    start = sessionnet_hoechste_id(http, basis, ext)
    if not start:
        raise RuntimeError("keine Sitzungsnummern gefunden (RSS/Kalender nicht erreichbar)")
    print(f"   SessionNet: starte bei Sitzung {start}, gehe rückwärts")
    alt_in_folge, gesehen = 0, set()
    for ksinr in range(start + 20, max(start + 20 - max_sitzungen, 0), -1):
        url = urljoin(basis, f"si0057.{ext}?__ksinr={ksinr}")
        try:
            html = http.get(url, cache=False)
        except PermissionError:
            raise
        except Exception:  # noqa: BLE001
            continue
        if "fehlermeldung" in html.lower()[:3000]:
            http.zaehle("Sitzung nicht vorhanden")
            continue
        soup = BeautifulSoup(html, "html.parser")
        seite = soup.get_text(" ", strip=True)
        datum = sitzungsdatum(seite)
        if seit and datum and datum < seit:
            alt_in_folge += 1
            if alt_in_folge >= stop_nach_alt:
                print(f"   SessionNet: {stop_nach_alt} Sitzungen in Folge vor {seit}, stoppe bei {ksinr}")
                return
            continue
        alt_in_folge = 0
        http.zaehle("Sitzungen geprüft")
        titel_el = soup.find("h1") or soup.find("title")
        gremium = titel_el.get_text(" ", strip=True) if titel_el else ""
        for a in soup.find_all("a", href=True):
            if not TOP_LINK_RE.search(a["href"]):
                continue
            zeile = a.find_parent("tr") or a.find_parent("li") or a.find_parent("div")
            kontext = (zeile.get_text(" ", strip=True) if zeile else a.get_text(" ", strip=True))[:400]
            if not PV_RE.search(kontext):
                continue
            ziel = urljoin(url, a["href"])
            if ziel in gesehen:
                continue
            gesehen.add(ziel)
            http.zaehle("PV-Tagesordnungspunkte")
            try:
                detail_html = http.get(ziel)
            except Exception:  # noqa: BLE001
                continue
            dsoup = BeautifulSoup(detail_html, "html.parser")
            # vom TOP zur Vorlage springen, wenn verlinkt
            vo = dsoup.find("a", href=re.compile(r"vo0050\.(?:php|asp)\?__kvonr=\d+"))
            if vo and "to0050" in ziel:
                vo_url = urljoin(ziel, vo["href"])
                if vo_url not in gesehen:
                    gesehen.add(vo_url)
                    try:
                        detail_html += "\n" + http.get(vo_url)
                        ziel = vo_url
                        dsoup = BeautifulSoup(detail_html, "html.parser")
                    except Exception:  # noqa: BLE001
                        pass
            pdfs = [urljoin(ziel, x["href"]) for x in dsoup.find_all("a", href=True)
                    if re.search(r"\.pdf|getfile", x["href"], re.I)][:2]
            anhang = "\n".join(dokument_text(http, u)[:15000] for u in pdfs)
            yield {"titel": kontext[:300], "datum": datum,
                   "text": f"{gremium}\n{kontext}\n{html_text(detail_html)}\n{anhang}",
                   "link": ziel, "dokument_link": pdfs[0] if pdfs else "",
                   "body_name": gremium or q.name}


def erstes_datum(text: str) -> str:
    for m in DATUM_RE.finditer(text):
        d, mo, y = map(int, m.groups())
        try:
            return dt.date(y, mo, d).isoformat()
        except ValueError:
            continue
    return ""


# --------------------------------------------------------------------------- #
# Extraktion
# --------------------------------------------------------------------------- #
def _zahl(s: str) -> float:
    return float(s.replace(".", "").replace(",", ".")) if s.count(",") == 1 or "." not in s else float(s)


_FIRMA_PRAEFIX = re.compile(
    r"^(?:(?:Antragsteller(?:in)?|Vorhabenträger(?:in)?|Investor(?:in)?|Projektentwickler(?:in)?|"
    r"Die|Der|Das|Den|Firma|Mit|Von|Durch|Ist|Hat|Haben)\s+)+")


def _firma_bereinigen(f: str) -> str:
    return _FIRMA_PRAEFIX.sub("", f.strip(" ,.:;"))


def extrahiere_heuristisch(roh: dict, q: Quelle) -> Beschluss:
    text = roh["text"]
    t_low = text.lower()
    titel_low = roh["titel"].lower()

    typ = ""
    for name, pat in reversed(BESCHLUSS_TYPEN):   # Titel zuerst, spätester Schritt gewinnt
        if re.search(pat, titel_low):
            typ = name
            break
    if not typ:
        for name, pat in BESCHLUSS_TYPEN:          # sonst erster Treffer im Text
            if re.search(pat, t_low):
                typ = name
                break

    konzept = next((n for n, p in KONZEPT_RE if re.search(p, t_low)), "")
    fl = FLAECHE_RE.search(text)
    lw = LEISTUNG_RE.search(text)
    firmen = [f.strip(" ,") for f in FIRMA_RE.findall(text) if "Amt" not in f.split()[0]]
    entw = ""
    m = re.search(r"vorhabenträger(?:in)?|projektentwickler(?:in)?|investor(?:in)?|antragsteller(?:in)?", text, re.I)
    if m:
        f = FIRMA_RE.search(text[m.start():m.start() + 250]) or FIRMA_RE.search(text[max(0, m.start() - 200):m.start()])
        entw = f.group(1) if f else ""
    entw = _firma_bereinigen(entw or (firmen[0] if firmen else ""))

    pn = PLANNAME_RE.search(roh["titel"]) or PLANNAME_RE.search(text)
    nr = PLANNR_RE.search(roh["titel"]) or PLANNR_RE.search(text[:2000])
    gm = GEMEINDE_RE.search(roh["titel"]) or GEMEINDE_RE.search(text[:3000])

    return Beschluss(
        quelle=q.name, landkreis=q.landkreis,
        gemeinde=gm.group(2).strip() if gm else roh.get("body_name", q.name),
        datum=roh["datum"], titel=roh["titel"][:300],
        beschlusstyp=typ or "Unklar", negativ=bool(NEGATIV_RE.search(t_low[:5000])),
        planname=pn.group(1).strip() if pn else "", plannr=nr.group(1) if nr else "",
        konzept=konzept,
        flaeche_ha=_zahl(fl.group(1)) if fl else None,
        leistung_mw=_zahl(lw.group(1)) if lw else None,
        entwickler=entw, link=roh["link"], dokument_link=roh.get("dokument_link", ""),
        bundesland=q.bundesland or "Brandenburg",
    )


GEMINI_PROMPT = """Du extrahierst Daten aus einer deutschen Ratsvorlage zu einem Photovoltaik-Bebauungsplan.
Antworte NUR mit JSON nach diesem Schema (null wenn unbekannt, nichts erfinden):
{"relevant": bool,  // false wenn es nicht um ein PV-Bauleitplanverfahren geht
 "gemeinde": str, "planname": str, "plannr": str,
 "beschlusstyp": one of ["Aufstellungsbeschluss","Frühzeitige Beteiligung","Entwurfs-/Auslegungsbeschluss","Abwägungsbeschluss","Satzungsbeschluss","Inkrafttreten/Bekanntmachung","Unklar"],
 "negativ": bool,   // abgelehnt, aufgehoben, Verfahren eingestellt
 "konzept": one of ["Freifläche","Agri-PV","Moor-PV","Floating-PV","Dach-PV",""],
 "flaeche_ha": number|null, "leistung_mw": number|null,
 "entwickler": str  // Vorhabenträger/Projektentwickler (Firma), sonst ""
}
TEXT:
"""


def extrahiere_gemini(roh: dict, q: Quelle, basis: Beschluss) -> Beschluss | None:
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        raise SystemExit("GEMINI_API_KEY ist nicht gesetzt.")
    model = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    body = {"contents": [{"parts": [{"text": GEMINI_PROMPT + roh["text"][:30000]}]}],
            "generationConfig": {"responseMimeType": "application/json", "temperature": 0}}
    for versuch in range(3):
        try:
            r = requests.post(url, params={"key": key}, json=body, timeout=90)
            r.raise_for_status()
            j = json.loads(r.json()["candidates"][0]["content"]["parts"][0]["text"])
            break
        except Exception as e:  # noqa: BLE001
            if versuch == 2:
                print(f"  ! Gemini-Fehler, nutze Heuristik: {e}", file=sys.stderr)
                return basis
            time.sleep(3 * (versuch + 1))
    if j.get("relevant") is False:
        return None
    for feld in ["gemeinde", "planname", "plannr", "beschlusstyp", "konzept", "entwickler"]:
        if j.get(feld):
            setattr(basis, feld, str(j[feld]))
    for feld in ["flaeche_ha", "leistung_mw"]:
        if isinstance(j.get(feld), (int, float)):
            setattr(basis, feld, float(j[feld]))
    if isinstance(j.get("negativ"), bool):
        basis.negativ = j["negativ"]
    basis.extraktion = "gemini"
    return basis


# --------------------------------------------------------------------------- #
# Projekt-Matching & Status
# --------------------------------------------------------------------------- #
def _norm(s: str) -> str:
    s = s.lower()
    s = re.sub(r"(vorhabenbezogene[rn]?|bebauungsplan(es)?|b-plan|vb-plan|nr\.?\s*\d+\w*|der gemeinde|gemeinde|stadt)", " ", s)
    return re.sub(r"[^a-zäöüß0-9]+", " ", s).strip()


def ordne_projekte_zu(beschluesse: list[Beschluss], schwelle: float = 0.82) -> None:
    projekte: list[tuple[str, str, str, str]] = []   # (id, gemeinde_norm, nr, name_norm)
    for b in sorted(beschluesse, key=lambda x: x.datum or "9999"):
        g, nr, nm = _norm(b.gemeinde), b.plannr.lower(), _norm(b.planname or b.titel)
        treffer = None
        for pid, pg, pnr, pnm in projekte:
            if pg != g:
                continue
            if nr and pnr and nr == pnr:
                treffer = pid
                break
            if nm and pnm and SequenceMatcher(None, nm, pnm).ratio() >= schwelle:
                treffer = pid
                break
        if not treffer:
            treffer = f"PV-{len(projekte) + 1:04d}"
            projekte.append((treffer, g, nr, nm))
        b.projekt_id = treffer


def status_fuer(events: list[Beschluss], heute: dt.date) -> str:
    events = sorted(events, key=lambda x: x.datum or "")
    letzte = events[-1]
    if any(e.negativ for e in events[-2:]):
        return "Abgelehnt"
    typen = {e.beschlusstyp for e in events}
    if typen & {"Satzungsbeschluss", "Inkrafttreten/Bekanntmachung"}:
        return "Genehmigt"
    if letzte.datum:
        d = dt.date.fromisoformat(letzte.datum)
        if (heute - d).days > RUHEND_NACH_MONATEN * 30:
            return "Ruhend"
    return "In Planung"


# --------------------------------------------------------------------------- #
# Excel-Export
# --------------------------------------------------------------------------- #
def schreibe_excel(beschluesse: list[Beschluss], quellen: list[Quelle], pfad: str,
                   protokoll: list[tuple[str, str, int]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    heute = dt.date.today()
    F = Font(name="Arial", size=10)
    FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    FL = Font(name="Arial", size=10, color="0563C1", underline="single")
    KOPF = PatternFill("solid", fgColor="1F4E78")
    STATUS_FARBE = {"In Planung": "FFF2CC", "Genehmigt": "E2EFDA", "Abgelehnt": "F8CBAD", "Ruhend": "D9D9D9"}

    wb = Workbook()

    def kopf(ws, spalten, breiten):
        ws.append(spalten)
        for i, c in enumerate(ws[1], 1):
            c.font, c.fill = FB, KOPF
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = breiten[i - 1]
        ws.freeze_panes = "A2"

    # --- Beschluesse
    wsb = wb.active
    wsb.title = "Beschluesse"
    sp_b = ["Projekt-ID", "Datum", "Gemeinde", "Landkreis", "Beschlusstyp", "Negativ",
            "Titel der Vorlage", "Link Vorlage", "Link Dokument", "Quelle (RIS)", "Extraktion"]
    kopf(wsb, sp_b, [11, 11, 22, 20, 26, 9, 70, 40, 40, 26, 11])
    for b in sorted(beschluesse, key=lambda x: (x.projekt_id, x.datum)):
        wsb.append([b.projekt_id, b.datum, b.gemeinde, b.landkreis, b.beschlusstyp,
                    "ja" if b.negativ else "", b.titel, b.link, b.dokument_link, b.quelle, b.extraktion])
        r = wsb.max_row
        for c in wsb[r]:
            c.font = F
            c.alignment = Alignment(vertical="top", wrap_text=c.column == 7)
        for col in (8, 9):
            c = wsb.cell(r, col)
            if c.value:
                c.hyperlink, c.font = c.value, FL
    wsb.auto_filter.ref = wsb.dimensions

    # --- Projekte
    wsp = wb.create_sheet("Projekte", 0)
    sp_p = ["Projekt-ID", "Gemeinde", "Landkreis", "Bundesland", "Planname", "Plan-Nr.",
            "Aufstellungskonzept", "Fläche (ha)", "Leistung (MW)", "Projektentwickler",
            "Status", "Letzter Beschluss", "Datum letzter Beschluss", "Erster Beschluss (Datum)",
            "Anzahl Beschlüsse", "Beschlusshistorie", "Links"]
    kopf(wsp, sp_p, [11, 22, 20, 13, 40, 9, 16, 11, 12, 32, 12, 26, 13, 13, 11, 60, 50])
    gruppen: dict[str, list[Beschluss]] = {}
    for b in beschluesse:
        gruppen.setdefault(b.projekt_id, []).append(b)

    def erstes(werte):
        return next((w for w in werte if w not in (None, "")), "")

    for pid in sorted(gruppen):
        ev = sorted(gruppen[pid], key=lambda x: x.datum or "")
        neu = list(reversed(ev))
        status = status_fuer(ev, heute)
        historie = "\n".join(f"{e.datum or 'o. D.'}: {e.beschlusstyp}{' (negativ)' if e.negativ else ''}" for e in ev)
        links = "\n".join(dict.fromkeys(x for e in ev for x in (e.link, e.dokument_link) if x))
        wsp.append([pid, ev[0].gemeinde, ev[0].landkreis, ev[0].bundesland,
                    erstes(e.planname for e in neu), erstes(e.plannr for e in neu),
                    erstes(e.konzept for e in neu), erstes(e.flaeche_ha for e in neu),
                    erstes(e.leistung_mw for e in neu), erstes(e.entwickler for e in neu),
                    status, ev[-1].beschlusstyp, ev[-1].datum, ev[0].datum, None, historie, links])
        r = wsp.max_row
        wsp.cell(r, 15).value = f'=COUNTIF(Beschluesse!$A:$A,A{r})'
        for c in wsp[r]:
            c.font = F
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (5, 10, 16, 17))
        wsp.cell(r, 11).fill = PatternFill("solid", fgColor=STATUS_FARBE.get(status, "FFFFFF"))
        wsp.cell(r, 8).number_format = "0.0"
        wsp.cell(r, 9).number_format = "0.0"
    wsp.auto_filter.ref = wsp.dimensions

    # --- Quellen
    wsq = wb.create_sheet("Quellen")
    kopf(wsq, ["Quelle", "Landkreis", "Typ", "URL", "Treffer", "Status Lauf"], [28, 20, 8, 70, 9, 50])
    for name, info, n in protokoll:
        q = next((x for x in quellen if x.name == name), None)
        wsq.append([name, q.landkreis if q else "", q.typ if q else "", q.url if q else "", n, info])
        for c in wsq[wsq.max_row]:
            c.font = F

    # --- Legende
    wsl = wb.create_sheet("Legende")
    zeilen = [
        ("Stand", heute.isoformat()),
        ("Projekte", "Eine Zeile pro Bebauungsplanverfahren; Werte stammen aus dem jüngsten Beschluss, der das Feld nennt."),
        ("Beschluesse", "Eine Zeile pro Ratsvorlage/Beschluss; Grundlage für Historie und Status."),
        ("Status", f"Abgelehnt = negativer Beschluss; Genehmigt = Satzungsbeschluss/Inkrafttreten; "
                   f"Ruhend = > {RUHEND_NACH_MONATEN} Monate kein Beschluss; sonst In Planung."),
        ("Extraktion", "heuristik = Regex-Regeln; gemini = LLM-Extraktion. Werte stichprobenartig prüfen."),
        ("Datum", "Datum der Vorlage laut RIS, nicht zwingend das Sitzungsdatum."),
    ]
    kopf(wsl, ["Feld", "Erklärung"], [16, 110])
    for z in zeilen:
        wsl.append(list(z))
        for c in wsl[wsl.max_row]:
            c.font = F
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(pfad)


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def lade_quellen(pfad: str) -> list[Quelle]:
    with open(pfad, encoding="utf-8") as f:
        return [Quelle(**{k: (v or "").strip() for k, v in row.items() if k in Quelle.__dataclass_fields__})
                for row in csv.DictReader(f) if row.get("url") and not row["name"].startswith("#")]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--quellen", default="quellen_brandenburg.csv")
    ap.add_argument("--seit", default="2022-01-01", help="nur Vorlagen ab diesem Datum (YYYY-MM-DD)")
    ap.add_argument("--llm", choices=["keins", "gemini"], default="keins")
    ap.add_argument("--out", default=f"PV_Beschluesse_Brandenburg_{dt.date.today():%Y-%m-%d}.xlsx")
    ap.add_argument("--json", help="zusätzlich Rohdaten als JSON speichern")
    ap.add_argument("--pause", type=float, default=REQUEST_PAUSE_S)
    a = ap.parse_args(argv)

    http = Http(pause=a.pause)
    quellen = lade_quellen(a.quellen)
    alle: list[Beschluss] = []
    protokoll = []

    for q in quellen:
        print(f"→ {q.name} ({q.typ})")
        try:
            typ, url = q.typ.lower(), q.url
            http.diag_reset()
            if typ == "auto":
                gefunden = oparl_entdecken(http, url)
                if gefunden:
                    typ, url = "oparl", gefunden
                elif ist_sessionnet(http, _basis(url)):
                    typ = "sessionnet"
                else:
                    typ = "html"
                print(f"   erkannt: {typ} {url if gefunden else ''}")
            qq = Quelle(q.name, q.landkreis, typ, url, q.bundesland)
            if typ == "oparl":
                roh_iter = oparl_vorlagen(http, qq, a.seit)
            elif typ == "sessionnet":
                roh_iter = sessionnet_vorlagen(http, qq, a.seit)
            else:
                roh_iter = html_vorlagen(http, qq, a.seit)
            n = 0
            for roh in roh_iter:
                b = extrahiere_heuristisch(roh, qq)
                if a.llm == "gemini":
                    b = extrahiere_gemini(roh, qq, b)
                    if b is None:
                        continue
                elif not PLANUNG_RE.search(roh["text"][:20000]):
                    continue
                alle.append(b)
                n += 1
                print(f"   + {b.datum or 'o. D.'} | {b.gemeinde} | {b.beschlusstyp} | {b.titel[:70]}")
            d = http.diag_reset()
            info = ", ".join(f"{k}: {v}" for k, v in sorted(d.items()))
            print(f"   = {n} Treffer | {info or 'keine Abrufe'}")
            hinweis = " – robots.txt verbietet Zugriff, Betreiber um Erlaubnis fragen" if d.get("robots.txt verbietet") and not d.get("Seiten geladen") else ""
            protokoll.append((q.name, f"ok ({typ}) | {info}{hinweis}"[:300], n))
        except Exception as e:  # noqa: BLE001
            d = http.diag_reset()
            info = ", ".join(f"{k}: {v}" for k, v in sorted(d.items()))
            print(f"   ! Fehler: {e} | {info}")
            protokoll.append((q.name, f"Fehler: {e} | {info}"[:300], 0))

    ordne_projekte_zu(alle)
    schreibe_excel(alle, quellen, a.out, protokoll)
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump([asdict(b) for b in alle], f, ensure_ascii=False, indent=2)
    print(f"\n{len(alle)} Beschlüsse, {len({b.projekt_id for b in alle})} Projekte → {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
